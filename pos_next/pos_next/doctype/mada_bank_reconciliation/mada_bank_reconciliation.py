# Copyright (c) 2026, MUBTKIR and contributors
# For license information, please see license.txt

import csv
import io
import re
from datetime import date, datetime

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import flt, getdate

# كلمات دالّة لاكتشاف الأعمدة (عربي/إنجليزي)
HEADER_KEYWORDS = {
    "date": ["date", "trans date", "transaction date", "posting date", "value date",
             "التاريخ", "تاريخ", "تاريخ العملية", "تاريخ الحركة"],
    "amount": ["amount", "credit", "deposit", "cr amount", "credit amount",
               "المبلغ", "مبلغ", "دائن", "ايداع", "إيداع", "القيمة", "قيمة", "مبلغ الإيداع"],
    "description": ["description", "desc", "details", "narration", "remarks", "notes", "particulars",
                    "البيان", "الوصف", "التفاصيل", "تفاصيل", "المرجع", "ملاحظات", "بيان"],
    "settlement_no": ["settlement", "settlement no", "batch", "رقم الموازنة",
                      "الموازنة", "موازنة", "reference no", "رقم المرجع"],
}

HEADER_HINTS = ["date", "amount", "credit", "التاريخ", "تاريخ", "المبلغ", "دائن", "البيان", "الوصف"]

ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")


class MadaBankReconciliation(Document):
    # ------------------------------------------------------------------ #
    #  رفع كشف البنك وتعبئة الصفوف تلقائيًا                                #
    # ------------------------------------------------------------------ #
    @frappe.whitelist()
    def parse_statement_file(self):
        if not self.bank_statement_file:
            frappe.throw(_("أرفق ملف كشف البنك أولًا (Excel أو CSV)."))

        rows = self._read_rows_from_file()
        if not rows:
            frappe.throw(_("لم أتمكن من قراءة أي صف إيداع من الملف. تأكد أن الملف Excel/CSV وفيه عمود مبلغ."))

        settings = frappe.get_cached_doc("Mada Reconciliation Settings")
        pattern = (settings.settlement_regex or "").strip()

        self.set("statement_rows", [])
        added = 0
        for r in rows:
            settlement_no = r.get("settlement_no")
            if not settlement_no and pattern and r.get("description"):
                try:
                    m = re.search(pattern, r["description"])
                    if m:
                        settlement_no = m.group(1) if m.groups() else m.group(0)
                except re.error:
                    pass

            self.append("statement_rows", {
                "txn_date": r.get("date"),
                "settlement_no": (settlement_no or "").strip() or None,
                "amount": r.get("amount"),
                "description": r.get("description"),
            })
            added += 1

        self.save()
        frappe.db.commit()
        return {"rows": added}

    def _read_rows_from_file(self):
        file_doc = frappe.get_doc("File", {"file_url": self.bank_statement_file})
        file_name = (file_doc.file_name or "").lower()

        if file_name.endswith(".csv") or file_name.endswith(".txt"):
            content = file_doc.get_content()
            if isinstance(content, bytes):
                content = content.decode("utf-8-sig", errors="ignore")
            table = [row for row in csv.reader(io.StringIO(content))]
        else:
            from openpyxl import load_workbook

            wb = load_workbook(file_doc.get_full_path(), read_only=True, data_only=True)
            ws = wb.active
            table = [list(row) for row in ws.iter_rows(values_only=True)]
            wb.close()

        return self._map_rows(table)

    def _map_rows(self, table):
        if not table:
            return []

        # اكتشاف صف العناوين خلال أول 15 صفًا
        header_idx = None
        for i, row in enumerate(table[:15]):
            cells = [self._norm(c) for c in row if c is not None]
            if any(any(hint in cell for cell in cells) for hint in HEADER_HINTS):
                header_idx = i
                break
        if header_idx is None:
            header_idx = 0

        header = [self._norm(c) for c in table[header_idx]]
        col = {}
        for key, kws in HEADER_KEYWORDS.items():
            for ci, h in enumerate(header):
                if h and any(kw in h for kw in kws):
                    col[key] = ci
                    break

        def cell(row, key):
            ci = col.get(key)
            if ci is None or ci >= len(row):
                return None
            return row[ci]

        rows = []
        for row in table[header_idx + 1:]:
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            amount = self._to_float(cell(row, "amount"))
            settlement_no = cell(row, "settlement_no")
            desc = cell(row, "description")

            # إيداعات فقط: تجاهل الصفوف بلا مبلغ موجب ما لم يوجد رقم موازنة
            if (amount is None or amount <= 0) and not settlement_no:
                continue

            rows.append({
                "date": self._to_date(cell(row, "date")),
                "amount": amount,
                "description": str(desc).strip() if desc is not None else None,
                "settlement_no": str(settlement_no).strip() if settlement_no not in (None, "") else None,
            })
        return rows

    @staticmethod
    def _norm(value):
        if value is None:
            return ""
        return str(value).strip().lower().translate(ARABIC_DIGITS)

    @staticmethod
    def _to_float(value):
        if value is None or value == "":
            return None
        if isinstance(value, (int, float)):
            return flt(value)
        text = str(value).translate(ARABIC_DIGITS)
        negative = "(" in text and ")" in text
        # فواصل عربية: ٬ (آلاف) و ٫ (عشري) و ، (آلاف)
        text = text.replace("٬", "").replace("،", "").replace(",", "")
        text = text.replace("٫", ".")
        text = re.sub(r"[^0-9.\-]", "", text)
        if text in ("", "-", "."):
            return None
        try:
            num = float(text)
        except ValueError:
            return None
        return -num if negative else num

    @staticmethod
    def _to_date(value):
        if value is None or value == "":
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip().translate(ARABIC_DIGITS)
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        try:
            return getdate(text)
        except Exception:
            return None

    # ------------------------------------------------------------------ #
    #  المطابقة                                                           #
    # ------------------------------------------------------------------ #
    @frappe.whitelist()
    def run_matching(self):
        """
        يطابق صفوف كشف البنك مع سجلات 'Mada Settlement' عبر رقم الموازنة:
          - رقم موجود في الطرفين + المبلغ متساوٍ  -> Matched
          - رقم موجود لكن المبلغ مختلف             -> Amount Mismatch
          - موازنة مسجّلة بلا صف بنكي               -> Not in Bank
          - صف بنكي بلا موازنة مسجّلة                -> Extra in Bank (عدّ فقط)
        """
        matched = not_in_bank = mismatch = extra = 0

        bank_by_no = {}
        for row in (self.statement_rows or []):
            key = (row.settlement_no or "").strip()
            if not key:
                continue
            bank_by_no.setdefault(key, []).append(row)

        filters = {}
        if self.from_date:
            filters["settlement_date"] = [">=", getdate(self.from_date)]
        settlements = frappe.get_all(
            "Mada Settlement",
            filters=filters,
            fields=["name", "settlement_no", "amount", "settlement_date"],
        )

        settlement_nos = set()
        for s in settlements:
            if self.to_date and getdate(s.settlement_date) > getdate(self.to_date):
                continue

            key = (s.settlement_no or "").strip()
            settlement_nos.add(key)
            matches = bank_by_no.get(key)

            if not matches:
                frappe.db.set_value("Mada Settlement", s.name, {
                    "status": "Not in Bank",
                    "bank_amount": 0,
                    "difference": -flt(s.amount),
                    "bank_matched_on": None,
                })
                not_in_bank += 1
                continue

            bank_amount = sum(flt(r.amount) for r in matches)
            diff = flt(bank_amount) - flt(s.amount)
            status = "Matched" if abs(diff) < 0.005 else "Amount Mismatch"

            frappe.db.set_value("Mada Settlement", s.name, {
                "status": status,
                "bank_amount": bank_amount,
                "difference": diff,
                "bank_matched_on": self.reconciliation_date,
            })
            if status == "Matched":
                matched += 1
            else:
                mismatch += 1

        for key, matches in bank_by_no.items():
            if key not in settlement_nos:
                extra += len(matches)

        self.db_set("matched_count", matched)
        self.db_set("not_in_bank_count", not_in_bank)
        self.db_set("mismatch_count", mismatch)
        self.db_set("extra_count", extra)
        frappe.db.commit()

        return {
            "matched": matched,
            "not_in_bank": not_in_bank,
            "mismatch": mismatch,
            "extra": extra,
        }
